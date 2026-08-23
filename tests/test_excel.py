"""Tests for the Excel + glossary layer.

The two exit-gate properties get the most attention:

* **Formatting survives.** A client's sheet must come back looking identical.
* **Nothing is written until accepted.** Extraction must never mutate the source.

Model-backed translation is covered in `test_translate.py` and skipped when the
weights are absent.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from backend import db
from backend.features import excel, glossary


def sheet_bytes() -> bytes:
    """A workbook with the things a real client sheet has."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Catalogue"

    ws["A1"] = "Product"
    ws["B1"] = "Unit"
    ws["C1"] = "Price"
    ws["D1"] = "Qty"
    ws["E1"] = "Total"
    for cell in ("A1", "B1", "C1", "D1", "E1"):
        ws[cell].font = Font(bold=True, size=13, color="FFFFFF")
        ws[cell].fill = PatternFill("solid", start_color="203864")
        ws[cell].alignment = Alignment(horizontal="center")
        ws[cell].border = Border(bottom=Side(style="medium"))

    rows = [
        ("Coconut oil", "1 litre bottle", 180.5, 3),
        ("Fresh mango pickle", "500 g jar", 95.0, 10),
        ("Cardamom", "100 g pack", 320.75, 2),
    ]
    for i, (name, unit, price, qty) in enumerate(rows, start=2):
        ws[f"A{i}"] = name
        ws[f"B{i}"] = unit
        ws[f"C{i}"] = price
        ws[f"C{i}"].number_format = "#,##0.00"
        ws[f"D{i}"] = qty
        ws[f"E{i}"] = f"=C{i}*D{i}"  # a formula, must be left alone

    ws["A6"] = "മൊത്തം"  # already Malayalam — nothing to do
    ws["A7"] = 12345  # a number
    ws.column_dimensions["A"].width = 28
    ws.merge_cells("A9:C9")
    ws["A9"] = "Contact our office for bulk orders"

    second = wb.create_sheet("Notes")
    second["A1"] = "Delivery within three days"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- extraction -----------------------------------------------------------


def test_extract_finds_the_english_cells() -> None:
    found = excel.extract(sheet_bytes())
    sources = {c.source for c in found.cells}
    assert "Coconut oil" in sources
    assert "Contact our office for bulk orders" in sources
    assert "Delivery within three days" in sources
    assert found.sheets == ["Catalogue", "Notes"]


def test_extract_skips_formulas() -> None:
    found = excel.extract(sheet_bytes())
    assert found.skipped_formulas == 3
    assert not any(c.source.startswith("=") for c in found.cells)


def test_extract_skips_numbers_and_existing_malayalam() -> None:
    found = excel.extract(sheet_bytes())
    sources = {c.source for c in found.cells}
    assert "മൊത്തം" not in sources, "already Malayalam — must not be re-translated"
    assert "12345" not in sources
    assert not any(c.source == "180.5" for c in found.cells)


def test_cell_keys_address_exactly() -> None:
    found = excel.extract(sheet_bytes())
    by_key = {c.key: c.source for c in found.cells}
    assert by_key["Catalogue!A2"] == "Coconut oil"
    assert by_key["Notes!A1"] == "Delivery within three days"


def test_extract_does_not_modify_the_source() -> None:
    """Exit gate: nothing is written until the operator accepts."""
    original = sheet_bytes()
    copy = bytes(original)
    excel.extract(original)
    assert original == copy


def test_unique_sources_deduplicates() -> None:
    wb = Workbook()
    ws = wb.active
    for row in range(1, 51):
        ws[f"A{row}"] = "Price per kilogram"  # a header repeated 50 times
        ws[f"B{row}"] = f"Item {row}"
    buffer = io.BytesIO()
    wb.save(buffer)
    found = excel.extract(buffer.getvalue())
    assert len(found.cells) == 100
    # 50 distinct "Item n" plus one repeated header.
    assert len(excel.unique_sources(found.cells)) == 51


def test_unreadable_file_is_rejected_clearly() -> None:
    with pytest.raises(excel.ExcelError, match="could not be read"):
        excel.extract(b"this is not a spreadsheet")


# --- applying -------------------------------------------------------------


def test_apply_preserves_formatting() -> None:
    """The whole reason we assign cell.value instead of rebuilding the sheet."""
    data = sheet_bytes()
    out = excel.apply(data, {"Catalogue!A2": "വെളിച്ചെണ്ണ"})
    ws = load_workbook(io.BytesIO(out))["Catalogue"]

    assert ws["A2"].value == "വെളിച്ചെണ്ണ"
    # Header styling intact
    assert ws["A1"].font.bold is True
    assert ws["A1"].font.size == 13
    assert ws["A1"].fill.start_color.rgb.endswith("203864")
    assert ws["A1"].alignment.horizontal == "center"
    # Column width, number format, merges intact
    assert ws.column_dimensions["A"].width == 28
    assert ws["C2"].number_format == "#,##0.00"
    assert "A9:C9" in [str(r) for r in ws.merged_cells.ranges]


def test_apply_leaves_formulas_alone() -> None:
    out = excel.apply(sheet_bytes(), {"Catalogue!E2": "should not land"})
    ws = load_workbook(io.BytesIO(out))["Catalogue"]
    assert ws["E2"].value == "=C2*D2"


def test_apply_leaves_numbers_alone() -> None:
    out = excel.apply(sheet_bytes(), {"Catalogue!C2": "nope"})
    ws = load_workbook(io.BytesIO(out))["Catalogue"]
    assert ws["C2"].value == 180.5


def test_apply_accepts_a_partial_review() -> None:
    """A half-reviewed sheet is a valid export; untouched cells stay English."""
    out = excel.apply(sheet_bytes(), {"Catalogue!A2": "വെളിച്ചെണ്ണ"})
    ws = load_workbook(io.BytesIO(out))["Catalogue"]
    assert ws["A2"].value == "വെളിച്ചെണ്ണ"
    assert ws["A3"].value == "Fresh mango pickle"


def test_apply_ignores_unknown_sheets() -> None:
    out = excel.apply(sheet_bytes(), {"NoSuchSheet!A1": "x"})
    assert load_workbook(io.BytesIO(out)).sheetnames == ["Catalogue", "Notes"]


# --- glossary -------------------------------------------------------------


def terms(**pairs: str) -> list[tuple[str, str]]:
    return glossary.compile_terms(
        [{"source_term": k.replace("_", " "), "target_term": v} for k, v in pairs.items()]
    )


def test_longest_term_wins() -> None:
    t = terms(oil="എണ്ണ", coconut_oil="വെളിച്ചെണ്ണ")
    assert [pair[0] for pair in t][0] == "coconut oil"
    assert glossary.apply_glossary_only("Coconut oil", t) == "വെളിച്ചെണ്ണ"


def test_terms_are_case_insensitive() -> None:
    t = terms(pickle="അച്ചാർ")
    assert glossary.apply_glossary_only("PICKLE", t) == "അച്ചാർ"
    assert glossary.apply_glossary_only("Pickle", t) == "അച്ചാർ"


def test_term_does_not_match_inside_a_word() -> None:
    t = terms(oil="എണ്ണ")
    masked = glossary.mask("The boiler is broken", t)
    assert masked.terms == {}, "'oil' must not match inside 'boiler'"


def test_fully_covered_cells_skip_the_model() -> None:
    t = terms(coconut_oil="വെളിച്ചെണ്ണ")
    assert glossary.is_fully_covered("Coconut oil", t) is True
    assert glossary.is_fully_covered("Coconut oil, 1 litre", t) is False


def test_lost_term_is_reported_not_silently_dropped() -> None:
    """A model that swallows a placeholder must not lose a brand name."""
    t = terms(coconut_oil="വെളിച്ചെണ്ണ")
    masked = glossary.mask("Coconut oil bottle", t)
    # Simulate the model dropping the placeholder entirely.
    restored = glossary.restore("കുപ്പി", masked)
    assert restored.lost == ["വെളിച്ചെണ്ണ"]
    assert "വെളിച്ചെണ്ണ" in restored.text


def test_client_product_code_is_never_deleted() -> None:
    """A source "X0X" looks like our placeholder. It must survive untouched."""
    t = terms(coconut_oil="വെളിച്ചെണ്ണ")
    masked = glossary.mask("Cable X0X coconut oil 5m", t)
    restored = glossary.restore(masked.text, masked)
    assert "X0X" in restored.text
    assert "വെളിച്ചെണ്ണ" in restored.text


def test_no_terms_is_a_passthrough() -> None:
    masked = glossary.mask("Anything at all", [])
    assert masked.text == "Anything at all"
    assert glossary.restore(masked.text, masked).text == "Anything at all"


# --- per-client scoping ---------------------------------------------------


def test_glossary_is_scoped_per_client() -> None:
    """Exit gate: two clients with a conflicting term must not interfere."""
    a = db.add_client("Scoped A")
    b = db.add_client("Scoped B")
    db.upsert_terms(a["id"], [{"source_term": "Fresh", "target_term": "ആംബർ ഫ്രഷ്"}])
    db.upsert_terms(b["id"], [{"source_term": "Fresh", "target_term": "പുതിയ"}])

    assert db.get_glossary(a["id"])[0]["target_term"] == "ആംബർ ഫ്രഷ്"
    assert db.get_glossary(b["id"])[0]["target_term"] == "പുതിയ"


def test_correcting_a_term_fixes_every_later_use() -> None:
    """Exit gate: the first correction is manual, the rest are automatic."""
    client = db.add_client("Correcting C")
    db.upsert_terms(client["id"], [{"source_term": "pickle", "target_term": "wrong"}])
    db.upsert_terms(client["id"], [{"source_term": "pickle", "target_term": "അച്ചാർ"}])

    rows = db.get_glossary(client["id"])
    assert len(rows) == 1, "correcting must update, not add a duplicate"
    assert rows[0]["target_term"] == "അച്ചാർ"

    t = glossary.compile_terms(rows)
    for sentence in ("Mango pickle", "Lime pickle", "PICKLE"):
        assert "അച്ചാർ" in glossary.apply_glossary_only(sentence, t)
