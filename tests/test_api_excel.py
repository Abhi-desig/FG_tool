"""API tests for Phase 3: clients, glossary, translate, export."""

from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from backend import db
from backend.features import translate, verify
from backend.main import app

# base_url is a loopback address deliberately: `guard.LocalOnlyMiddleware`
# refuses any Host that is not this machine, and TestClient's default
# `http://testserver` is exactly the DNS-rebinding shape it exists to stop.
client = TestClient(app, base_url="http://127.0.0.1:8000")

HAS_ENGINE = translate.available_engine() is not None
needs_engine = pytest.mark.skipif(not HAS_ENGINE, reason="no engine downloaded")


def sheet() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Catalogue"
    ws["A1"] = "Product"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "Coconut oil"
    ws["A3"] = "Fresh mango pickle"
    ws["B2"] = 180.5
    ws["C2"] = "=B2*2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def upload() -> dict[str, tuple[str, bytes, str]]:
    return {
        "file": (
            "catalogue.xlsx",
            sheet(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


def wait_for(job_id: str, timeout: float = 600.0) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        body = client.get(f"/api/jobs/{job_id}").json()
        if body["status"] not in {"queued", "running"}:
            return body
        time.sleep(0.1)
    raise AssertionError(f"job {job_id} did not finish")


@pytest.fixture
def a_client() -> int:
    made = client.post("/api/clients", json={"name": f"Test {time.time_ns()}"}).json()
    return int(made["id"])


# --- clients and glossary -------------------------------------------------


def clients(include_archived: bool = False) -> list[dict]:
    """Client rows out of the wrapped envelope.

    These routes returned bare lists while every other route in the app used an
    envelope (NEXT.md 3.12). Now `{"clients": [...], "count": n}`.
    """
    query = "?include_archived=true" if include_archived else ""
    return client.get(f"/api/clients{query}").json()["clients"]


def terms_of(client_id: int) -> list[dict]:
    return client.get(f"/api/clients/{client_id}/glossary").json()["terms"]


def test_client_lifecycle(a_client: int) -> None:
    assert a_client in [c["id"] for c in clients()]
    client.post(f"/api/clients/{a_client}/archive")
    assert a_client not in [c["id"] for c in clients()]
    assert a_client in [c["id"] for c in clients(include_archived=True)]


def test_client_list_is_wrapped_in_an_envelope() -> None:
    """The convention every other route in this app already followed."""
    body = client.get("/api/clients").json()
    assert isinstance(body, dict)
    assert isinstance(body["clients"], list)
    assert body["count"] == len(body["clients"])


def test_client_name_is_required() -> None:
    assert client.post("/api/clients", json={"name": ""}).status_code == 422


def test_glossary_round_trip(a_client: int) -> None:
    body = client.put(
        f"/api/clients/{a_client}/glossary",
        json={"terms": [{"source_term": "coconut oil", "target_term": "വെളിച്ചെണ്ണ"}]},
    ).json()
    assert body["terms"][0]["target_term"] == "വെളിച്ചെണ്ണ"
    assert body["client_id"] == a_client

    term_id = body["terms"][0]["id"]
    after = client.delete(f"/api/clients/{a_client}/glossary/{term_id}").json()
    assert after["terms"] == []
    assert after["count"] == 0


def test_glossary_put_still_accepts_a_bare_list(a_client: int) -> None:
    """The shipped `dist` on the shop PC may be older than this server."""
    body = client.put(
        f"/api/clients/{a_client}/glossary",
        json=[{"source_term": "pickle", "target_term": "അച്ചാർ"}],
    ).json()
    assert body["terms"][0]["target_term"] == "അച്ചാർ"


def test_correcting_a_term_updates_not_duplicates(a_client: int) -> None:
    """Exit gate: correcting once fixes every later occurrence."""
    client.put(
        f"/api/clients/{a_client}/glossary",
        json=[{"source_term": "pickle", "target_term": "wrong"}],
    )
    rows = client.put(
        f"/api/clients/{a_client}/glossary",
        json=[{"source_term": "pickle", "target_term": "അച്ചാർ"}],
    ).json()["terms"]
    assert len(rows) == 1
    assert rows[0]["target_term"] == "അച്ചാർ"


def test_two_clients_do_not_share_terms() -> None:
    """Exit gate: client A's brand terms must not leak into client B."""
    a = client.post("/api/clients", json={"name": f"A {time.time_ns()}"}).json()["id"]
    b = client.post("/api/clients", json={"name": f"B {time.time_ns()}"}).json()["id"]
    client.put(
        f"/api/clients/{a}/glossary",
        json=[{"source_term": "Fresh", "target_term": "ആംബർ ഫ്രഷ്"}],
    )
    client.put(
        f"/api/clients/{b}/glossary",
        json=[{"source_term": "Fresh", "target_term": "പുതിയ"}],
    )
    assert terms_of(a)[0]["target_term"] == "ആംബർ ഫ്രഷ്"
    assert terms_of(b)[0]["target_term"] == "പുതിയ"


def test_glossary_for_unknown_client_is_404() -> None:
    assert client.get("/api/clients/999999/glossary").status_code == 404


# --- inspect --------------------------------------------------------------


def test_inspect_counts_what_matters() -> None:
    body = client.post("/api/excel/inspect", files=upload()).json()
    assert body["sheets"] == ["Catalogue"]
    assert body["translatable"] == 3  # Product, Coconut oil, Fresh mango pickle
    assert body["skipped_formulas"] == 1
    assert body["skipped_non_text"] == 1  # the number


def test_inspect_rejects_a_non_spreadsheet() -> None:
    r = client.post(
        "/api/excel/inspect", files={"file": ("x.xlsx", b"not a sheet", "application/x")}
    )
    assert r.status_code == 400


def test_inspect_rejects_empty_upload() -> None:
    r = client.post("/api/excel/inspect", files={"file": ("x.xlsx", b"", "application/x")})
    assert r.status_code == 400


def test_translation_settings_lists_engines() -> None:
    body = client.get("/api/settings/translation").json()
    assert {e["key"] for e in body["engines"]} == {
        "opus-mt-en-ml",
        "indictrans2-en-indic",
    }


# --- translate and export -------------------------------------------------


def test_export_before_translation_is_rejected() -> None:
    r = client.post(
        "/api/excel/export", json={"job_id": "nope", "translations": {"A!A1": "x"}}
    )
    assert r.status_code == 409


@needs_engine
def test_translate_returns_a_review_grid_without_writing_a_file(a_client: int) -> None:
    """Exit gate: nothing is written until the operator accepts."""
    client.put(
        f"/api/clients/{a_client}/glossary",
        json=[{"source_term": "coconut oil", "target_term": "വെളിച്ചെണ്ണ"}],
    )
    started = client.post(
        "/api/excel/translate", files=upload(), data={"client_id": str(a_client)}
    ).json()
    done = wait_for(started["id"])
    assert done["status"] == "done", done

    rows = {r["key"]: r for r in done["result"]["rows"]}
    assert set(rows) == {"Catalogue!A1", "Catalogue!A2", "Catalogue!A3"}
    assert rows["Catalogue!A2"]["glossary_only"] is True
    assert rows["Catalogue!A2"]["translation"] == "വെളിച്ചെണ്ണ"
    assert done["result"]["from_glossary"] >= 1
    # No file is offered by the translate step itself.
    assert "translated.xlsx" not in str(done["result"])


@needs_engine
def test_export_applies_operator_edits_and_keeps_formatting(a_client: int) -> None:
    started = client.post(
        "/api/excel/translate", files=upload(), data={"client_id": str(a_client)}
    ).json()
    done = wait_for(started["id"])
    assert done["status"] == "done", done

    # The operator overrides one row by hand — that edit must win.
    edits = {r["key"]: r["translation"] for r in done["result"]["rows"]}
    edits["Catalogue!A3"] = "എന്റെ സ്വന്തം വിവർത്തനം"

    got = client.post(
        "/api/excel/export", json={"job_id": started["id"], "translations": edits}
    )
    assert got.status_code == 200
    ws = load_workbook(io.BytesIO(got.content))["Catalogue"]
    assert ws["A3"].value == "എന്റെ സ്വന്തം വിവർത്തനം"
    assert ws["A1"].font.bold is True, "formatting must survive"
    assert ws["C2"].value == "=B2*2", "formulas must survive"
    assert ws["B2"].value == 180.5, "numbers must survive"


@needs_engine
def test_translating_without_a_client_still_works() -> None:
    """No glossary chosen is valid — it just means no forced terms."""
    started = client.post("/api/excel/translate", files=upload()).json()
    done = wait_for(started["id"])
    assert done["status"] == "done", done
    assert all(not r["glossary_only"] for r in done["result"]["rows"])


def test_translate_with_unknown_client_is_404() -> None:
    r = client.post(
        "/api/excel/translate", files=upload(), data={"client_id": "999999"}
    )
    assert r.status_code == 404


def test_translate_rejects_a_sheet_with_nothing_to_do() -> None:
    wb = Workbook()
    wb.active["A1"] = 42
    buffer = io.BytesIO()
    wb.save(buffer)
    r = client.post(
        "/api/excel/translate",
        files={"file": ("nums.xlsx", buffer.getvalue(), "application/x")},
    )
    assert r.status_code == 400
    assert "needs translating" in r.json()["detail"]


# --- the optional Claude check ---------------------------------------------
#
# No test here makes a real API call. What is worth testing at this level is the
# wiring: that a correction reaches the right cell of the review grid, that the
# offline text is kept beside it, and — above all — that asking for the check is
# a deliberate act. The check itself is tested in `tests/test_verify.py`.


def test_inspect_quotes_the_check_without_running_it() -> None:
    body = client.post("/api/excel/inspect", files=upload()).json()
    quote = body["verify"]
    assert quote["is_estimate"] is True
    assert quote["rows"] >= 1
    assert quote["cost_paise"] > 0
    # Quoting is free and must not need a key.
    assert "configured" in quote


@needs_engine
def test_the_check_is_off_unless_asked_for(monkeypatch: pytest.MonkeyPatch) -> None:
    """It spends money and sends the sheet off the machine. Never by default."""
    called = {"n": 0}

    def spy(*args: object, **kwargs: object) -> None:
        called["n"] += 1
        raise AssertionError("the paid check ran without being asked for")

    monkeypatch.setattr(verify, "check_rows", spy)

    started = client.post("/api/excel/translate", files=upload()).json()
    done = wait_for(started["id"])

    assert done["status"] == "done", done
    assert called["n"] == 0
    assert done["result"]["review"] == {"requested": False}


@needs_engine
def test_a_correction_reaches_the_right_cell(monkeypatch: pytest.MonkeyPatch) -> None:
    """The wiring that matters: verdicts are zipped back onto the right rows."""

    def fake_check(
        pairs: list[tuple[str, str]], *args: object, **kwargs: object
    ) -> verify.VerifyResult:
        # Correct exactly one source, leave the rest, so a misalignment shows.
        verdicts = [
            verify.Verdict("തേങ്ങാ എണ്ണ", checked=True, corrected=True)
            if source == "Coconut oil"
            else verify.Verdict(current, checked=True)
            for source, current in pairs
        ]
        return verify.VerifyResult(
            verdicts=verdicts, model="fake-model", cost_paise=250, checked=len(pairs),
            corrected=1,
        )

    monkeypatch.setattr(verify, "check_rows", fake_check)

    started = client.post(
        "/api/excel/translate", files=upload(), data={"check_with_claude": "true"}
    ).json()
    done = wait_for(started["id"])
    assert done["status"] == "done", done

    rows = {r["key"]: r for r in done["result"]["rows"]}
    corrected = rows["Catalogue!A2"]
    assert corrected["source"] == "Coconut oil"
    assert corrected["translation"] == "തേങ്ങാ എണ്ണ"
    assert corrected["verify_corrected"] is True
    # The offline reading is kept, not discarded — the operator decides.
    assert corrected["offline_translation"] != "തേങ്ങാ എണ്ണ"

    # Every other row is untouched and not falsely marked as changed.
    for key, row in rows.items():
        if key == "Catalogue!A2":
            continue
        assert row["verify_corrected"] is False
        assert row["translation"] == row["offline_translation"]

    review = done["result"]["review"]
    assert review["requested"] is True
    assert review["model"] == "fake-model"
    assert review["cost_rupees"] == 2.5
    assert done["result"]["verify_corrected"] == 1


@needs_engine
def test_a_failed_check_still_returns_a_usable_grid(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A refused key must not cost the operator the translation they paid for."""

    def fake_check(
        pairs: list[tuple[str, str]], *args: object, **kwargs: object
    ) -> verify.VerifyResult:
        return verify.VerifyResult(
            verdicts=[verify.Verdict(current) for _, current in pairs],
            model="fake-model",
            error="That API key was rejected.",
            unchecked=len(pairs),
        )

    monkeypatch.setattr(verify, "check_rows", fake_check)

    started = client.post(
        "/api/excel/translate", files=upload(), data={"check_with_claude": "true"}
    ).json()
    done = wait_for(started["id"])

    assert done["status"] == "done", done
    assert done["result"]["review"]["error"] == "That API key was rejected."
    assert len(done["result"]["rows"]) == 3
    assert all(r["checked"] is False for r in done["result"]["rows"])


# --- the corrections memory ------------------------------------------------
#
# ADR-029. The promise is "correct it once and the same spelling mistake cannot
# happen again", so these test the whole loop: an edit is learned on export, and
# the next sheet is filled from it without the model or the paid check running.


def _corrections(client_id: int | None = None) -> list[dict]:
    params = {} if client_id is None else {"client_id": client_id}
    return client.get("/api/corrections", params=params).json()["corrections"]


@pytest.fixture(autouse=True)
def empty_memory() -> None:
    """Each test starts with nothing remembered.

    The memory is shop-wide by design, so without this one test's export would
    silently satisfy the next test's assertions.
    """
    db.connect().execute("DELETE FROM corrections")
    db.connect().commit()


def _translate_and_export(edits: dict[str, str], **export_body: object) -> dict:
    started = client.post("/api/excel/translate", files=upload()).json()
    done = wait_for(started["id"])
    assert done["status"] == "done", done
    final = {r["key"]: r["translation"] for r in done["result"]["rows"]}
    final.update(edits)
    response = client.post(
        "/api/excel/export",
        json={"job_id": started["id"], "translations": final, **export_body},
    )
    assert response.status_code == 200, response.text
    return {"job": done, "response": response}


@needs_engine
def test_an_edited_cell_is_remembered_after_export() -> None:
    """The operator fixes a cell once; that is the whole feature.

    Only the edited cell is learned — the rows they left alone are the model's
    unreviewed output and must not become permanent answers.
    """
    started = client.post("/api/excel/translate", files=upload()).json()
    done = wait_for(started["id"])
    rows = {r["source"]: r for r in done["result"]["rows"]}

    final = {r["key"]: r["translation"] for r in done["result"]["rows"]}
    final[rows["Coconut oil"]["key"]] = "വെളിച്ചെണ്ണ"
    response = client.post(
        "/api/excel/export", json={"job_id": started["id"], "translations": final}
    )
    assert response.status_code == 200
    assert response.headers["X-Corrections-Remembered"] == "1"

    remembered = {c["source"]: c["target"] for c in _corrections()}
    assert remembered == {"Coconut oil": "വെളിച്ചെണ്ണ"}


@needs_engine
def test_an_untouched_machine_translation_is_never_remembered() -> None:
    """The review grid exists to distrust the model. Learning its raw output
    would turn one unreviewed guess into a permanent answer."""
    result = _translate_and_export({})
    assert result["response"].headers["X-Corrections-Remembered"] == "0"
    assert _corrections() == []


@needs_engine
def test_exporting_without_remember_learns_nothing() -> None:
    started = client.post("/api/excel/translate", files=upload()).json()
    done = wait_for(started["id"])
    final = {r["key"]: "എന്തെങ്കിലും" for r in done["result"]["rows"]}
    response = client.post(
        "/api/excel/export",
        json={"job_id": started["id"], "translations": final, "remember": False},
    )
    assert response.status_code == 200
    assert response.headers["X-Corrections-Remembered"] == "0"
    assert _corrections() == []


@needs_engine
def test_a_remembered_cell_comes_back_filled_on_the_next_sheet() -> None:
    """The payoff, end to end."""
    client.put(
        "/api/corrections",
        json={"corrections": [{"source": "Coconut oil", "target": "ഓർത്തുവച്ചത്"}]},
    )
    started = client.post("/api/excel/translate", files=upload()).json()
    done = wait_for(started["id"])
    rows = {r["source"]: r for r in done["result"]["rows"]}

    assert rows["Coconut oil"]["translation"] == "ഓർത്തുവച്ചത്"
    assert rows["Coconut oil"]["from_memory"] is True
    assert rows["Coconut oil"]["warnings"] == [], "an approved cell must not be flagged"
    assert done["result"]["from_memory"] == 1
    # Everything else still went through the model.
    assert rows["Fresh mango pickle"]["from_memory"] is False


def test_a_remembered_cell_is_never_sent_to_claude(monkeypatch: pytest.MonkeyPatch) -> None:
    """Paying a model to second-guess the operator's own approved answer is
    both a waste and a way to lose an answer that was already right."""
    client.put(
        "/api/corrections",
        json={"corrections": [{"source": "Coconut oil", "target": "ഓർത്തുവച്ചത്"}]},
    )
    sent: list[list[tuple[str, str]]] = []

    def fake_check(pairs, *args: object, **kwargs: object) -> verify.VerifyResult:
        sent.append(list(pairs))
        return verify.VerifyResult(
            verdicts=[verify.Verdict(c, checked=True) for _, c in pairs],
            model="fake",
            checked=len(pairs),
        )

    monkeypatch.setattr(verify, "check_rows", fake_check)
    monkeypatch.setattr(
        translate,
        "translate_rows",
        lambda sources, terms, *a, **k: _stub_rows(sources, k.get("memory") or {}),
    )

    started = client.post(
        "/api/excel/translate", files=upload(), data={"check_with_claude": "true"}
    ).json()
    done = wait_for(started["id"])
    assert done["status"] == "done", done

    assert sent, "the check should still have run for the unremembered rows"
    assert "Coconut oil" not in [source for source, _ in sent[0]]
    assert done["result"]["review"]["skipped_exact"] >= 1


def _stub_rows(sources: list[str], memory: dict[str, str]) -> list[translate.Row]:
    """Offline translation without a model, honouring the memory."""
    from backend import textkey

    rows = []
    for source in sources:
        remembered = memory.get(textkey.normalise(source))
        rows.append(
            translate.Row(
                source=source,
                translation=remembered or f"[{source}]",
                from_memory=bool(remembered),
            )
        )
    return rows


def test_a_glossary_only_cell_is_never_sent_to_claude(
    monkeypatch: pytest.MonkeyPatch, a_client: int
) -> None:
    """Defect: every row went to the paid check, including the ones
    `translate.py` calls exact by construction. A client's approved glossary
    term could be silently replaced by the model that was meant to check it."""
    client.put(
        f"/api/clients/{a_client}/glossary",
        json={"terms": [{"source_term": "Coconut oil", "target_term": "വെളിച്ചെണ്ണ"}]},
    )
    sent: list[list[tuple[str, str]]] = []

    def fake_check(pairs, *args: object, **kwargs: object) -> verify.VerifyResult:
        sent.append(list(pairs))
        return verify.VerifyResult(
            verdicts=[verify.Verdict("REPLACED", checked=True, corrected=True) for _ in pairs],
            model="fake",
            checked=len(pairs),
            corrected=len(pairs),
        )

    monkeypatch.setattr(verify, "check_rows", fake_check)
    monkeypatch.setattr(
        translate,
        "translate_rows",
        lambda sources, terms, *a, **k: [
            translate.Row(
                source=s,
                translation="വെളിച്ചെണ്ണ" if s == "Coconut oil" else f"[{s}]",
                glossary_only=(s == "Coconut oil"),
            )
            for s in sources
        ],
    )

    started = client.post(
        "/api/excel/translate",
        files=upload(),
        data={"client_id": str(a_client), "check_with_claude": "true"},
    ).json()
    done = wait_for(started["id"])
    rows = {r["source"]: r for r in done["result"]["rows"]}

    assert "Coconut oil" not in [source for source, _ in sent[0]]
    assert rows["Coconut oil"]["translation"] == "വെളിച്ചെണ്ണ", (
        "the paid check overwrote a term the client had approved"
    )
    assert rows["Coconut oil"]["verify_corrected"] is False


def test_the_quote_does_not_count_cells_the_memory_already_covers() -> None:
    """Defect: `inspect` took no client, so the estimate counted rows that
    would never reach the model — an over-quote on the one number the operator
    uses to decide whether to spend."""
    before = client.post("/api/excel/inspect", files=upload()).json()
    client.put(
        "/api/corrections",
        json={"corrections": [{"source": "Coconut oil", "target": "വെളിച്ചെണ്ണ"}]},
    )
    after = client.post("/api/excel/inspect", files=upload()).json()

    assert after["from_memory"] == 1
    assert after["to_translate"] == before["to_translate"] - 1
    assert after["verify"]["rows"] == before["verify"]["rows"] - 1
    assert after["verify"]["cost_paise"] < before["verify"]["cost_paise"]


def test_a_clients_correction_overrides_the_shop_wide_one(a_client: int) -> None:
    client.put(
        "/api/corrections",
        json={"corrections": [{"source": "Coconut oil", "target": "SHOP WIDE"}]},
    )
    client.put(
        "/api/corrections",
        json={
            "client_id": a_client,
            "corrections": [{"source": "Coconut oil", "target": "THIS CLIENT"}],
        },
    )
    body = client.post(
        "/api/excel/inspect", files=upload(), data={"client_id": str(a_client)}
    ).json()
    assert body["from_memory"] == 1

    from backend import db as _db

    assert _db.corrections_map(a_client)["coconut oil"] == "THIS CLIENT"
    assert _db.corrections_map(None)["coconut oil"] == "SHOP WIDE"


def test_another_clients_sheet_still_gets_the_shop_wide_correction(a_client: int) -> None:
    """A house name written by sound is right whoever the client is."""
    client.put(
        "/api/corrections",
        json={"corrections": [{"source": "Coconut oil", "target": "SHOP WIDE"}]},
    )
    other = client.post("/api/clients", json={"name": "Someone else"}).json()["id"]
    from backend import db as _db

    assert _db.corrections_map(other)["coconut oil"] == "SHOP WIDE"


def test_correcting_a_remembered_cell_replaces_it_rather_than_duplicating_it() -> None:
    for target in ("FIRST", "SECOND"):
        client.put(
            "/api/corrections",
            json={"corrections": [{"source": "Coconut oil", "target": target}]},
        )
    rows = _corrections()
    assert len(rows) == 1
    assert rows[0]["target"] == "SECOND"


# --- the spreadsheet round trip --------------------------------------------


def _pairs_file(pairs: list[tuple[str, str]]) -> dict:
    from backend.features import excel as excel_feature

    return {
        "file": (
            "corrections.xlsx",
            excel_feature.write_pairs(pairs, scope="All clients"),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


def test_loading_a_two_column_sheet_fills_the_memory() -> None:
    body = client.post(
        "/api/corrections/import",
        files=_pairs_file([("ELAVUNKAL VEEDU", "എളവുങ്കൽ വീട്"), ("Guardian", "രക്ഷിതാവ്")]),
    ).json()
    assert body["added"] == 2
    assert {c["source"] for c in _corrections()} == {"ELAVUNKAL VEEDU", "Guardian"}


def test_a_swapped_column_import_is_refused() -> None:
    """A reversed memory does not sit there looking wrong — it silently fills
    the next sheet's cells with English."""
    response = client.post(
        "/api/corrections/import",
        files=_pairs_file([("എളവുങ്കൽ വീട്", "ELAVUNKAL VEEDU")]),
    )
    assert response.status_code == 400
    assert "wrong way round" in response.json()["detail"]
    assert _corrections() == []


def test_downloaded_corrections_load_back_unchanged() -> None:
    client.put(
        "/api/corrections",
        json={"corrections": [{"source": "Guardian", "target": "രക്ഷിതാവ്"}]},
    )
    downloaded = client.get("/api/corrections/export")
    assert downloaded.status_code == 200

    db.connect().execute("DELETE FROM corrections")
    db.connect().commit()
    body = client.post(
        "/api/corrections/import",
        files={
            "file": (
                "corrections.xlsx",
                downloaded.content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    ).json()
    assert body["added"] == 1
    assert {c["source"]: c["target"] for c in _corrections()} == {"Guardian": "രക്ഷിതാവ്"}


def test_forgetting_a_job_removes_only_what_it_taught() -> None:
    db.upsert_corrections(None, [("Kept", "സൂക്ഷിച്ചു")], learned_from="other-job")
    db.upsert_corrections(None, [("Gone", "പോയി")], learned_from="bad-job")
    body = client.post("/api/corrections/forget-job", json={"job_id": "bad-job"}).json()
    assert body["forgotten"] == 1
    assert {c["source"] for c in _corrections()} == {"Kept"}


def test_a_correction_can_be_removed_one_at_a_time() -> None:
    client.put(
        "/api/corrections",
        json={"corrections": [{"source": "Guardian", "target": "രക്ഷിതാവ്"}]},
    )
    row_id = _corrections()[0]["id"]
    assert client.delete(f"/api/corrections/{row_id}").status_code == 200
    assert _corrections() == []


def test_an_untouched_remembered_row_is_not_relearned() -> None:
    """Re-saving it looks harmless — the value is identical — but it re-stamps
    the entry with this job's id, so a later forget-job would delete a
    correction the operator made weeks ago on a different sheet."""
    db.upsert_corrections(
        None, [("Coconut oil", "വെളിച്ചെണ്ണ")], learned_from="an-older-job"
    )
    started = client.post("/api/excel/translate", files=upload()).json()
    done = wait_for(started["id"])
    final = {r["key"]: r["translation"] for r in done["result"]["rows"]}
    response = client.post(
        "/api/excel/export", json={"job_id": started["id"], "translations": final}
    )

    assert response.status_code == 200
    assert response.headers["X-Corrections-Remembered"] == "0"
    kept = _corrections()[0]
    assert kept["learned_from"] == "an-older-job", (
        "an untouched memory row was re-stamped with the new job"
    )
